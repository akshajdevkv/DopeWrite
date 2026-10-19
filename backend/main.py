"""DopeWrite extraction API. Documents are processed in memory and never saved."""
import asyncio
import io
import os
import zipfile
from pathlib import Path

from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from backend.auth import verify_request
from pptx import Presentation
from backend.topics import Lesson, organize, build_lesson, source_blocks, local_plan
from backend.pdf_text import extract_pdf, tesseract_command

load_dotenv(Path(__file__).resolve().parent.parent / '.env.local')
load_dotenv(Path(__file__).resolve().parent.parent / '.env')
load_dotenv(Path(__file__).with_name('.env'))
app = FastAPI(title='DopeWrite API')
app.add_middleware(CORSMiddleware, allow_origins=[os.getenv('FRONTEND_ORIGIN', 'http://localhost:3000')], allow_methods=['GET', 'POST'], allow_headers=['Content-Type'])
app.middleware('http')(verify_request)
MAX_BYTES = 15 * 1024 * 1024
MAX_TEXT = 45000

def document_text(data: bytes, extension: str) -> str:
    if extension == '.pdf':
        source = extract_pdf(data, MAX_TEXT)
    else:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 60 * 1024 * 1024:
                    raise ValueError('Expanded file too large')
            presentation = Presentation(io.BytesIO(data))
            slides = []
            for i, slide in enumerate(presentation.slides):
                texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
                for shape in slide.shapes:
                    if shape.has_table:
                        texts.extend('\n'.join(cell.text for cell in row.cells) for row in shape.table.rows)
                if any(shape.shape_type == 13 for shape in slide.shapes):
                    texts.append(f'[Diagram/image on slide {i + 1}; view the original presentation]')
                text = '\n'.join(texts).strip()
                if text:
                    slides.append(text)
            source = '\n\n'.join(slides)
        except Exception:
            raise HTTPException(422, 'Could not read this presentation. Upload a valid PowerPoint (.pptx) file.')
    if not source.strip():
        raise HTTPException(422, 'No readable text was found. Try a clearer PDF scan or a presentation with selectable text.')
    if len(source) > MAX_TEXT:
        raise HTTPException(413, 'This document has too much text for one session. Split it into smaller files.')
    return source


def create_lesson(source: str, title: str) -> Lesson:
    blocks = source_blocks(source)
    return build_lesson(blocks, local_plan(blocks), title)


@app.get('/health')
def health():
    return {'status': 'ok', 'provider': 'groq' if os.getenv('GROQ_API_KEY') else 'local', 'extraction_configured': True, 'ocr_configured': bool(tesseract_command()), 'verbatim': True}


@app.post('/api/extract', response_model=Lesson)
async def extract(file: UploadFile = File(...)):
    extension = Path(file.filename or '').suffix.lower()
    if extension not in {'.pdf', '.pptx'}:
        raise HTTPException(415, 'Upload a PDF or PowerPoint (.pptx) file.')
    data = await file.read(MAX_BYTES + 1)
    await file.close()
    if not data or len(data) > MAX_BYTES:
        raise HTTPException(413, 'Choose a non-empty file smaller than 15 MB.')
    source = await asyncio.to_thread(document_text, data, extension)
    blocks, plan, organization, notice = await organize(source)
    return build_lesson(blocks, plan, Path(file.filename or 'Document').stem, organization, notice)


from backend.evaluation import EvaluationRequest, Evaluation, evaluate_submission


@app.post('/api/evaluate', response_model=Evaluation)
async def evaluate(submission: EvaluationRequest):
    return await evaluate_submission(submission)


from backend.processing import ProcessingRequest, ProcessingResponse, edit_processing


@app.post('/api/processing-chat', response_model=ProcessingResponse)
async def processing_chat(request: ProcessingRequest):
    return await edit_processing(request)


from backend.recall import RecallRequest, QuestionSet, frame_questions


@app.post('/api/recall-questions', response_model=QuestionSet)
async def recall_questions(request: RecallRequest):
    return await frame_questions(request)


from backend.review import ReviewedDocument, prepare_review, reviewed_lesson


@app.post('/api/prepare', response_model=ReviewedDocument)
async def prepare(file: UploadFile = File(...)):
    extension = Path(file.filename or '').suffix.lower()
    if extension not in {'.pdf', '.pptx'}:
        raise HTTPException(415, 'Upload a PDF or PowerPoint (.pptx) file.')
    data = await file.read(MAX_BYTES + 1)
    await file.close()
    if not data or len(data) > MAX_BYTES:
        raise HTTPException(413, 'Choose a non-empty file smaller than 15 MB.')
    source = await asyncio.to_thread(document_text, data, extension)
    return prepare_review(source, Path(file.filename or 'Document').stem)


@app.post('/api/review', response_model=Lesson)
async def review(document: ReviewedDocument):
    return reviewed_lesson(document)
