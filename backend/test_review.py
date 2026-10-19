import io
import unittest
from unittest.mock import AsyncMock,patch

from fastapi.testclient import TestClient
from pptx import Presentation
from backend.main import app


class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.client=TestClient(app)

    def test_upload_prepares_editable_document_without_ai_or_practice(self):
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='Memory';slide.placeholders[1].text='This wording needs correcting.'
        data=io.BytesIO();deck.save(data)
        with patch('httpx.AsyncClient.post',new_callable=AsyncMock) as provider:
            response=self.client.post('/api/prepare',files={'file':('lecture.pptx',data.getvalue())})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['title'],'lecture')
        self.assertIn('This wording needs correcting.',[block['text'] for block in response.json()['blocks']])
        self.assertNotIn('sections',response.json())
        provider.assert_not_called()

    def test_manual_headings_and_corrections_are_authoritative(self):
        document={'title':'My corrected notes','blocks':[
            {'kind':'heading','heading_level':1,'text':'First topic'},
            {'kind':'body','text':'The user corrected these EXACT words.'},
            {'kind':'equation','text':'E = mc²'},
            {'kind':'heading','heading_level':2,'text':'A custom subtopic'},
            {'kind':'body','text':'My second corrected paragraph stays unchanged.'},
            {'kind':'ignore','text':'Do not practice this footer.'}]}
        with patch('httpx.AsyncClient.post',new_callable=AsyncMock) as provider:
            response=self.client.post('/api/review',json=document)
        self.assertEqual(response.status_code,200)
        lesson=response.json()
        self.assertEqual(lesson['organization'],'manual')
        self.assertEqual([section['title'] for section in lesson['sections']],['First topic','A custom subtopic'])
        self.assertEqual([section['text'] for section in lesson['sections']],['The user corrected these EXACT words.','My second corrected paragraph stays unchanged.'])
        self.assertEqual(lesson['source_blocks'][1],'The user corrected these EXACT words.')
        provider.assert_not_called()

    def test_normal_text_is_not_reclassified_by_heuristics(self):
        response=self.client.post('/api/review',json={'title':'Reviewed','blocks':[{'kind':'body','text':'NOT A HEADING'},{'kind':'body','text':'The user explicitly marked this as normal text.'}]})
        self.assertEqual(response.status_code,200)
        self.assertIn('NOT A HEADING',response.json()['sections'][0]['text'])

    def test_references_only_document_is_editable_but_cannot_start_practice(self):
        response=self.client.post('/api/review',json={'title':'Notes','blocks':[{'kind':'heading','text':'Only a heading'}]})
        self.assertEqual(response.status_code,422)
        self.assertIn('Normal text',response.json()['detail'])

    def test_invalid_document_is_rejected(self):
        for document in [{'title':'','blocks':[{'text':'Words','kind':'body'}]},
                         {'title':'Notes','blocks':[{'text':'Words','kind':'unexpected'}]},
                         {'title':'Notes','blocks':[{'text':'x'*45001,'kind':'body'}]}]:
            with self.subTest(document=document['title']):
                self.assertEqual(self.client.post('/api/review',json=document).status_code,422)

    def test_image_markers_cannot_accidentally_become_typing(self):
        response=self.client.post('/api/review',json={'title':'Notes','blocks':[{'text':'[Diagram/image on PDF page 1]','kind':'body'}]})
        self.assertEqual(response.status_code,422)

if __name__=='__main__':
    unittest.main()
